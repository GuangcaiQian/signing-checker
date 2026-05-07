# -*- coding: utf-8 -*-
import os, sys, json, re, base64, io, threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from concurrent.futures import ThreadPoolExecutor, as_completed

import fitz
import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

CN = ("SimSun", 12)
EN = ("Times New Roman", 12)
CN_TITLE = ("SimSun", 20, "bold")
CN_BOLD = ("SimSun", 12, "bold")
EN_LOG = ("Times New Roman", 11)

CFG = os.path.join(os.path.dirname(os.path.abspath(sys.argv[0])), "config.json")

def load_cfg():
    try:
        with open(CFG, "r", encoding="utf-8") as f: return json.load(f)
    except: return {}

def save_cfg(c):
    try:
        with open(CFG, "w", encoding="utf-8") as f: json.dump(c, f, ensure_ascii=False)
    except: pass

def check_one(args):
    pdf_path, idx, api_url, api_key, model, prompt, cancel_event = args
    if cancel_event.is_set():
        return idx + 1, None
    doc = fitz.open(pdf_path)
    page = doc[idx]
    pix = page.get_pixmap(matrix=fitz.Matrix(0.75, 0.75))
    img_bytes = pix.tobytes("png")
    doc.close()
    if cancel_event.is_set():
        return idx + 1, None
    b64 = base64.standard_b64encode(img_bytes).decode()
    r = requests.post(api_url,
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        json={"model": model, "messages": [{"role":"user","content":[
            {"type":"image_url","image_url":{"url":f"data:image/png;base64,{b64}"}},
            {"type":"text","text":prompt}
        ]}], "max_tokens":300}, timeout=90)
    text = r.json()["choices"][0]["message"]["content"]
    m = re.search(r"\{[^}]+\}", text)
    data = json.loads(m.group()) if m else {"类型":"未知","质量员":"解析失败","机组长":"解析失败","监理工程师":"解析失败"}
    return idx + 1, data

class App:
    def __init__(self, root):
        self.root = root
        self.root.title("焊接记录签字校核工具")
        self.root.geometry("700x560")
        self.root.configure(bg="#f8f9fa")
        self.root.resizable(False, False)
        self.running = False
        self.cancel = threading.Event()
        self.total = 0
        self.results = {}
        cfg = load_cfg()
        self._build(cfg)

    def _build(self, cfg):
        BG = "#f8f9fa"; ACCENT = "#4361ee"; FG = "#2b2d42"
        PAD = dict(padx=16, pady=6)
        style = ttk.Style()
        style.theme_use("clam")
        style.configure("TLabel", background=BG, foreground=FG, font=CN)
        style.configure("TButton", font=CN)
        style.configure("Big.TButton", font=CN_BOLD, padding=10)
        style.configure("Cancel.TButton", font=CN_BOLD, padding=10)
        style.configure("TEntry", font=EN, padding=4)
        style.configure("TLabelframe", background=BG)
        style.configure("TLabelframe.Label", background=BG, foreground=FG, font=CN_BOLD)
        style.configure("Horizontal.TProgressbar", troughcolor="#e0e0e0", background=ACCENT)
        style.configure("TSpinbox", font=EN)

        # 标题
        tf = tk.Frame(self.root, bg=BG)
        tf.pack(pady=(16, 4))
        tk.Label(tf, text="焊", font=CN_TITLE, bg=BG, fg=ACCENT).pack(side="left")
        tk.Label(tf, text="接", font=CN_TITLE, bg=BG, fg=ACCENT).pack(side="left")
        tk.Label(tf, text="记", font=CN_TITLE, bg=BG, fg=ACCENT).pack(side="left")
        tk.Label(tf, text="录", font=CN_TITLE, bg=BG, fg=ACCENT).pack(side="left")
        tk.Label(tf, text="签", font=CN_TITLE, bg=BG, fg=ACCENT).pack(side="left")
        tk.Label(tf, text="字", font=CN_TITLE, bg=BG, fg=ACCENT).pack(side="left")
        tk.Label(tf, text="校", font=CN_TITLE, bg=BG, fg=ACCENT).pack(side="left")
        tk.Label(tf, text="核", font=CN_TITLE, bg=BG, fg=ACCENT).pack(side="left")

        # 文件
        f1 = ttk.LabelFrame(self.root, text="  文件  ", padding=10)
        f1.pack(fill="x", **PAD)
        r1 = ttk.Frame(f1); r1.pack(fill="x")
        tk.Label(r1, text="PDF", font=EN, bg=BG).pack(side="left")
        tk.Label(r1, text="：", font=CN, bg=BG).pack(side="left")
        self.v_pdf = tk.StringVar()
        ttk.Entry(r1, textvariable=self.v_pdf, width=42).pack(side="left", padx=4)
        ttk.Button(r1, text="选择文件", command=self._sel_pdf).pack(side="left")
        r2 = ttk.Frame(f1); r2.pack(fill="x", pady=(6,0))
        tk.Label(r2, text="输出", font=EN, bg=BG).pack(side="left")
        tk.Label(r2, text="：", font=CN, bg=BG).pack(side="left")
        self.v_out = tk.StringVar(value=cfg.get("out",""))
        ttk.Entry(r2, textvariable=self.v_out, width=42).pack(side="left", padx=4)
        ttk.Button(r2, text="选择目录", command=self._sel_out).pack(side="left")

        # API
        f2 = ttk.LabelFrame(self.root, text="  API  ", padding=10)
        f2.pack(fill="x", **PAD)
        tk.Label(f2, text="地址：", font=CN, bg=BG).grid(row=0, column=0, sticky="w")
        self.v_url = tk.StringVar(value=cfg.get("url","https://dashscope.aliyuncs.com/compatible-mode/v1/chat/completions"))
        ttk.Entry(f2, textvariable=self.v_url, width=55).grid(row=0, column=1, columnspan=3, padx=4)
        tk.Label(f2, text="Key：", font=CN, bg=BG).grid(row=1, column=0, sticky="w", pady=(6,0))
        self.v_key = tk.StringVar(value=cfg.get("key",""))
        ttk.Entry(f2, textvariable=self.v_key, width=55, show="*").grid(row=1, column=1, columnspan=3, padx=4, pady=(6,0))
        tk.Label(f2, text="模型：", font=CN, bg=BG).grid(row=2, column=0, sticky="w", pady=(6,0))
        self.v_model = tk.StringVar(value=cfg.get("model","qwen-vl-turbo"))
        ttk.Entry(f2, textvariable=self.v_model, width=22).grid(row=2, column=1, sticky="w", padx=4, pady=(6,0))
        tk.Label(f2, text="并发：", font=CN, bg=BG).grid(row=2, column=2, sticky="e", pady=(6,0))
        self.v_workers = tk.IntVar(value=cfg.get("workers",5))
        ttk.Spinbox(f2, from_=1, to=20, textvariable=self.v_workers, width=4).grid(row=2, column=3, sticky="w", pady=(6,0))

        # 按钮 + 进度
        f3 = ttk.Frame(self.root); f3.pack(fill="x", **PAD)
        self.btn = ttk.Button(f3, text="  开始校核  ", style="Big.TButton", command=self._go)
        self.btn.pack(side="left")
        self.btn_cancel = ttk.Button(f3, text="  取消  ", style="Cancel.TButton", command=self._cancel)
        self.v_prog = tk.StringVar(value="就绪")
        tk.Label(f3, textvariable=self.v_prog, font=EN).pack(side="left", padx=16)
        self.pbar = ttk.Progressbar(f3, length=340, mode="determinate")
        self.pbar.pack(side="left", padx=8)

        # 日志
        f4 = ttk.LabelFrame(self.root, text="  日志  ", padding=4)
        f4.pack(fill="both", expand=True, **PAD)
        self.log = tk.Text(f4, state="disabled", bg="white", fg=FG,
                           font=EN_LOG, relief="flat", bd=2)
        sb = ttk.Scrollbar(f4, command=self.log.yview)
        self.log.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        self.log.pack(fill="both", expand=True)

    def _sel_pdf(self):
        p = filedialog.askopenfilename(filetypes=[("PDF","*.pdf")])
        if p:
            self.v_pdf.set(p)
            if not self.v_out.get(): self.v_out.set(os.path.dirname(p))

    def _sel_out(self):
        p = filedialog.askdirectory()
        if p: self.v_out.set(p)

    def _wlog(self, msg):
        self.log.configure(state="normal")
        self.log.insert("end", msg+"\n")
        self.log.see("end")
        self.log.configure(state="disabled")

    def _go(self):
        pdf = self.v_pdf.get().strip()
        out = self.v_out.get().strip()
        url = self.v_url.get().strip()
        key = self.v_key.get().strip()
        if not pdf or not os.path.exists(pdf):
            messagebox.showerror("错误","请选择PDF文件"); return
        if not url or not key:
            messagebox.showerror("错误","请填写API地址和Key"); return
        if not out: out = os.path.dirname(pdf); self.v_out.set(out)
        save_cfg({"url":url,"key":key,"model":self.v_model.get(),"out":out,"workers":self.v_workers.get()})
        self.btn.configure(state="disabled")
        self.cancel.clear()
        self.btn_cancel.pack(side="left", padx=(12,0))
        self.running = True
        threading.Thread(target=self._run, args=(pdf,out,url,key), daemon=True).start()

    def _cancel(self):
        self.cancel.set()
        self.btn_cancel.configure(state="disabled")
        self._wlog("  正在取消...")

    def _run(self, pdf, out, url, key):
        model = self.v_model.get().strip()
        workers = self.v_workers.get()
        prompt = ('这是一张焊接记录表扫描页。请先判断表格类型：'
                  '如果是"管道组对焊接记录"，检查底部"机组长"和"质量员"签字位；'
                  '如果是"焊口返修记录"，检查底部"质量员"和"监理工程师"签字位。'
                  '很淡的笔迹也算有，完全空白才算无。'
                  '回答JSON格式：{"类型":"组对/返修","质量员":"有/无","机组长":"有/无","监理工程师":"有/无"}'
                  '组对记录只需填质量员和机组长，返修记录只需填质量员和监理工程师，不涉及的填"无"。')
        try:
            doc = fitz.open(pdf)
            total = doc.page_count
            doc.close()
        except Exception as e:
            self.root.after(0, lambda: messagebox.showerror("错误", f"无法打开PDF: {e}"))
            self.root.after(0, self._reset)
            return
        self.total = total
        self.root.after(0, lambda: self.pbar.configure(maximum=total, value=0))
        self.root.after(0, self._wlog, f"共 {total} 页，开始检测...")
        self.results = {}
        miss_zl, miss_jz, miss_jl, errs = [], [], [], []
        done = 0
        pool = ThreadPoolExecutor(max_workers=workers)
        try:
            futs = {pool.submit(check_one,(pdf,i,url,key,model,prompt,self.cancel)):i for i in range(total)}
            for f in as_completed(futs):
                if self.cancel.is_set():
                    for ff in futs:
                        ff.cancel()
                    break
                done += 1
                try:
                    pn, d = f.result()
                    if d is None:
                        continue
                    self.results[pn] = d
                    tp = d.get("类型","?"); zl = d.get("质量员","?"); jz = d.get("机组长","?"); jl = d.get("监理工程师","?")
                    missing = []
                    if tp == "组对":
                        if jz == "无": miss_jz.append(pn); missing.append(f"机组长={jz}")
                        if zl == "无": miss_zl.append(pn); missing.append(f"质量员={zl}")
                    elif tp == "返修":
                        if jl == "无": miss_jl.append(pn); missing.append(f"监理工程师={jl}")
                        if zl == "无": miss_zl.append(pn); missing.append(f"质量员={zl}")
                    if missing:
                        self.root.after(0, self._wlog, f"  第{pn}页 [{tp}]: {' '.join(missing)}  <---")
                    if done%50==0:
                        self.root.after(0, self._wlog, f"  进度 {done}/{total}")
                except Exception as e:
                    pn = futs[f]+1
                    self.results[pn] = {"类型":"错误","质量员":"错误","机组长":"错误","监理工程师":"错误"}
                    errs.append(pn)
                    self.root.after(0, self._wlog, f"  第{pn}页出错: {str(e)[:50]}")
                self.root.after(0, lambda d=done,t=total: (self.pbar.configure(value=d), self.v_prog.set(f"{d}/{t}")))
        finally:
            pool.shutdown(wait=False, cancel_futures=True)
        if self.cancel.is_set():
            self.root.after(0, self._wlog, f"\n  已取消，已完成 {done}/{total} 页\n")
            self.root.after(0, self._reset)
            self.running = False
            return
        xlsx = os.path.join(out, "签字校核结果.xlsx")
        self._xlsx(xlsx, miss_zl, miss_jz, miss_jl)
        miss_zl.sort(); miss_jz.sort(); miss_jl.sort()
        s = (f"\n================================\n"
             f"  校核完成\n"
             f"  总页数: {total}\n"
             f"  质量员缺签: {len(miss_zl)}页  {miss_zl or ''}\n"
             f"  机组长缺签: {len(miss_jz)}页  {miss_jz or ''}\n"
             f"  监理工程师缺签: {len(miss_jl)}页  {miss_jl or ''}\n"
             f"  错误: {len(errs)}页\n"
             f"  Excel: {xlsx}\n")
        self.root.after(0, self._wlog, s)
        self.root.after(0, self._reset)
        self.root.after(0, lambda: messagebox.showinfo("完成",
            f"质量员缺签: {len(miss_zl)}页\n机组长缺签: {len(miss_jz)}页\n监理工程师缺签: {len(miss_jl)}页"))
        self.running = False

    def _reset(self):
        self.btn.configure(state="normal")
        self.btn_cancel.pack_forget()
        self.btn_cancel.configure(state="normal")

    def _xlsx(self, path, miss_zl, miss_jz, miss_jl):
        wb = Workbook(); ws = wb.active; ws.title = "签字校核"
        hf = Font(bold=True, size=13, color="FFFFFF")
        hfl = PatternFill(start_color="4361ee", end_color="4361ee", fill_type="solid")
        red = PatternFill(start_color="ffc8c8", end_color="ffc8c8", fill_type="solid")
        grn = PatternFill(start_color="c8f7c5", end_color="c8f7c5", fill_type="solid")
        bdr = Border(Side("thin"),Side("thin"),Side("thin"),Side("thin"))
        ca = Alignment(horizontal="center")
        for i,h in enumerate(["页码","类型","质量员","机组长","监理工程师"],1):
            cl = ws.cell(row=1,column=i,value=h); cl.font=hf; cl.fill=hfl; cl.alignment=ca; cl.border=bdr
        for pn in range(1, self.total+1):
            r = self.results.get(pn,{})
            tp = r.get("类型","?"); zl = r.get("质量员","?"); jz = r.get("机组长","?"); jl = r.get("监理工程师","?")
            ws.cell(row=pn+1,column=1,value=pn).border=bdr; ws.cell(row=pn+1,column=1).alignment=ca
            c2 = ws.cell(row=pn+1,column=2,value=tp); c2.border=bdr; c2.alignment=ca
            c3 = ws.cell(row=pn+1,column=3,value=zl); c3.border=bdr; c3.alignment=ca; c3.fill=red if zl=="无" else grn
            c4 = ws.cell(row=pn+1,column=4,value=jz); c4.border=bdr; c4.alignment=ca; c4.fill=red if tp=="组对" and jz=="无" else (grn if tp=="组对" else PatternFill())
            c5 = ws.cell(row=pn+1,column=5,value=jl); c5.border=bdr; c5.alignment=ca; c5.fill=red if tp=="返修" and jl=="无" else (grn if tp=="返修" else PatternFill())
        ws.column_dimensions["A"].width=8; ws.column_dimensions["B"].width=10
        ws.column_dimensions["C"].width=12; ws.column_dimensions["D"].width=12; ws.column_dimensions["E"].width=16
        ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:E{self.total+1}"
        wb.save(path)

def main():
    root = tk.Tk()
    App(root)
    root.mainloop()

if __name__ == "__main__":
    main()
